#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Rebuild docs/解决方案型测试授权需求方案.docx by strictly mirroring 
'/Users/qianjincheng/Library/Mobile Documents/com~apple~CloudDocs/Sangfor/原型/采购取消/需求文档/采购转单需求文档（终稿）.docx'
1-to-1 in styling, fonts, tables, heading hierarchy, footers, and margins.
"""

import sys
import os
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

TEMPLATE_PATH = "/Users/qianjincheng/Library/Mobile Documents/com~apple~CloudDocs/Sangfor/原型/采购取消/需求文档/采购转单需求文档（终稿）.docx"
OUTPUT_DOCX_PATH = Path(__file__).resolve().parents[1] / "docs" / "解决方案型测试授权需求方案.docx"
OUTPUT_MD_PATH = Path(__file__).resolve().parents[1] / "docs" / "解决方案型测试授权需求方案.md"

def set_run_font(run, size=11.5, bold=False, color="000000", ea_font="宋体-简", ascii_font="宋体-简"):
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), ea_font)
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:cs"), ascii_font)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

def add_p(doc, text="", style="Normal", sa=4.0, sb=0.0, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    if sa is not None:
        p.paragraph_format.space_after = Pt(sa)
    if sb is not None:
        p.paragraph_format.space_before = Pt(sb)
    if line_spacing is not None:
        p.paragraph_format.line_spacing = line_spacing
    if text:
        run = p.add_run(text)
        size = 11.5
        if style == "Title":
            size = 22.0
        elif style == "Heading 1":
            size = 16.0
        elif style == "Heading 2":
            size = 13.0
        elif style == "Heading 3":
            size = 11.5
        set_run_font(run, size=size)
    return p

def add_subpart_header(doc, title):
    """Adds the five-part section header like '1、初始状态' with bold=True, size=12pt"""
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_before = Pt(4.0)
    p.paragraph_format.space_after = Pt(2.0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(title)
    set_run_font(run, size=12.0, bold=True, color="000000")
    return p

def add_body_p(doc, text, sa=4.0):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_before = Pt(0.0)
    p.paragraph_format.space_after = Pt(sa)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=11.5, bold=False, color="000000")
    return p

def add_caption(doc, text):
    """Caption for figures like '图 1  当前原型：...'"""
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_before = Pt(4.0)
    p.paragraph_format.space_after = Pt(6.0)
    run = p.add_run(text)
    set_run_font(run, size=10.0, bold=False, color="595959")
    return p

def set_cell_properties(cell, width_dxa=None, fill_hex=None):
    tcPr = cell._tc.get_or_add_tcPr()
    if width_dxa:
        tcW = tcPr.find(qn("w:tcW"))
        if tcW is None:
            tcW = OxmlElement("w:tcW")
            tcPr.append(tcW)
        tcW.set(qn("w:w"), str(width_dxa))
        tcW.set(qn("w:type"), "dxa")
        
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:color="CBCDD1" w:sz="4" w:space="0"/>
            <w:left w:val="single" w:color="CBCDD1" w:sz="4" w:space="0"/>
            <w:bottom w:val="single" w:color="CBCDD1" w:sz="4" w:space="0"/>
            <w:right w:val="single" w:color="CBCDD1" w:sz="4" w:space="0"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    if fill_hex:
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill_hex}"/>')
        tcPr.append(shd)
    
    vAlign = OxmlElement("w:vAlign")
    vAlign.set(qn("w:val"), "center")
    tcPr.append(vAlign)

def set_table_properties(table):
    tblPr = table._tbl.tblPr
    tblPr.append(parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:color="CBCDD1" w:sz="4" w:space="0"/>
            <w:left w:val="single" w:color="CBCDD1" w:sz="4" w:space="0"/>
            <w:bottom w:val="single" w:color="CBCDD1" w:sz="4" w:space="0"/>
            <w:right w:val="single" w:color="CBCDD1" w:sz="4" w:space="0"/>
            <w:insideH w:val="single" w:color="CBCDD1" w:sz="4" w:space="0"/>
            <w:insideV w:val="single" w:color="CBCDD1" w:sz="4" w:space="0"/>
        </w:tblBorders>
    '''))
    tblCellMar = parse_xml(f'''
        <w:tblCellMar {nsdecls("w")}>
            <w:top w:w="0" w:type="dxa"/>
            <w:left w:w="108" w:type="dxa"/>
            <w:bottom w:w="0" w:type="dxa"/>
            <w:right w:w="108" w:type="dxa"/>
        </w:tblCellMar>
    ''')
    tblPr.append(tblCellMar)
    tblPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))

def create_styled_table(doc, headers, data, col_widths_in=None, header_fill="E5F6FF", header_text_color="000000"):
    num_rows = len(data) + 1
    num_cols = len(headers)
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_properties(table)
    
    col_dxas = [int(w * 1440) for w in col_widths_in] if col_widths_in else [int(7.0 / num_cols * 1440)] * num_cols

    # Header row
    hdr_row = table.rows[0]
    for c_idx, h_text in enumerate(headers):
        cell = hdr_row.cells[c_idx]
        set_cell_properties(cell, width_dxa=col_dxas[c_idx], fill_hex=header_fill)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2.0)
        p.paragraph_format.space_after = Pt(2.0)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(h_text)
        set_run_font(run, size=11.5, bold=True, color=header_text_color)

    # Data rows
    for r_idx, row_data in enumerate(data):
        row = table.rows[r_idx + 1]
        for c_idx in range(num_cols):
            cell = row.cells[c_idx]
            val = row_data[c_idx] if c_idx < len(row_data) else ""
            set_cell_properties(cell, width_dxa=col_dxas[c_idx], fill_hex=None)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2.0)
            p.paragraph_format.space_after = Pt(2.0)
            p.paragraph_format.line_spacing = 1.15
            sub_lines = val.split("\n")
            for s_idx, sl in enumerate(sub_lines):
                if s_idx > 0:
                    p = cell.add_paragraph()
                    p.paragraph_format.space_before = Pt(1.0)
                    p.paragraph_format.space_after = Pt(1.0)
                    p.paragraph_format.line_spacing = 1.15
                run = p.add_run(sl)
                set_run_font(run, size=11.0, bold=False, color="000000")

    # Empty paragraph after table
    p_after = doc.add_paragraph(style="Normal")
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4.0)
    return table

def main():
    print(f"Loading template from {TEMPLATE_PATH}...")
    doc = Document(TEMPLATE_PATH)
    
    # Preserve sectPr and footers
    sectPr = doc.sections[0]._sectPr
    footer = doc.sections[0].footer
    for p in footer.paragraphs:
        p.text = "解决方案型测试授权需求文档"
        if p.runs:
            set_run_font(p.runs[0], size=9.0, bold=False, color="595959", ea_font="Microsoft YaHei", ascii_font="Microsoft YaHei")

    # Clear existing body children except sectPr
    body = doc._body._element
    for child in list(body):
        if child != sectPr:
            body.remove(child)

    print("Building Document Content matching 采购转单需求文档（终稿） 1-to-1...")

    # Title
    add_p(doc, "解决方案型测试授权需求文档（终稿）", style="Title", sa=8.0, sb=0.0)

    # Table 0: 版本变更记录
    t0_headers = ["时间", "变更人", "变更原因", "变更内容"]
    t0_data = [
        ["2026 年 8 月 6 日", "钱金城", "创建需求文档初稿", "梳理解决方案型测试授权申请概要方案、系统边界及整单审批原则。"],
        ["2026 年 8 月 19 日", "钱金城", "补充设备与 XaaS 并存规则及云图身份校验", "支持设备产品线与 SaaS-XDR 混排并存，补充客户云图账号/ID强制校验。"],
        ["2026 年 8 月 26 日", "钱金城", "落地统一授权天数模型与抽屉遮挡修复", "引入 authDays 动态计算到期日，解决配置抽屉与模块导航浮层遮挡。"],
        ["2026 年 9 月 15 日", "钱金城", "编写字段填写与审批操作说明", "明确借测单关联、各产品线参数卡片、审批结论与授权记录展示口径。"],
        ["2026 年 9 月 20 日", "钱金城", "完善五段式规范与业务方灵活配置建议落地", "对齐标准五段式结构；落地业务方关于业务申请信息与产品组合免代码灵活配置机制。"],
        ["2026 年 9 月 20 日", "钱金城", "精简 4.2/4.3 章节层级并强化 XaaS 云图 ID 规范", "消除冗余微观子节；业务信息直接对齐《安全方案测试申请规则.xlsx》动态渲染；重点强化 XaaS 产品线云图 ID 强阻断校验。"]
    ]
    create_styled_table(doc, t0_headers, t0_data, col_widths_in=[1.5, 0.85, 2.1, 3.0], header_fill="C7ECFF", header_text_color="000000")

    # 1. 需求背景
    add_p(doc, "1. 需求背景", style="Heading 1", sb=24.0, sa=4.0)
    add_p(doc, "1.1 业务现状", style="Heading 2", sb=10.0, sa=2.0)
    add_body_p(doc, "解决方案型测试授权通常由多个协同运作的产品线打包组成（如企业云一体化方案、分布式安全运营方案、全球加速及全球组网方案等）。当前业务中，因资产借测与测试授权系统割裂，业务人员需按单产品反复发起多次授权申请，流程繁琐且难以保证同一方案下各产品授权期限与客户环境的一致性。")
    add_body_p(doc, "同时，部分解决方案存在设备硬件产品线与 XaaS 云服务产品线混排的场景，不同产品线的授权对象和校验逻辑不同；新资产系统完成借测审批后，若缺乏统一关联机制，容易导致方案产品组合脱节、二次手工录入繁琐及授权结果无法原子交付的问题。")

    add_p(doc, "1.2 需求目标", style="Heading 2", sb=10.0, sa=2.0)
    add_body_p(doc, "在 W3 测试授权平台内完成解决方案配置维护、关联借测单发起、纯软件申请、多产品线工作台参数配置、整单审批与单一汇总授权记录展示。")
    add_body_p(doc, "系统支持一张申请单承接多个产品线（设备线与 XaaS 线并存），方案范围只读锁定，统一授权到期时间，全产品原子化开通与异常回滚；并响应业务方建议，实现业务申请信息与产品授权组合的灵活配置，避免方案推广时频繁进行代码开发。")

    # 2. 需求范围
    add_p(doc, "2. 需求范围", style="Heading 1", sb=24.0, sa=4.0)
    add_p(doc, "2.1 需求概述", style="Heading 2", sb=10.0, sa=2.0)
    
    t1_headers = ["系统", "模块", "需求概述", "说明"]
    t1_data = [
        ["W3", "解决方案配置抽屉", "维护解决方案编码、名称、产品线组合、授权天数及启停状态", "配置归属 W3，向新资产系统提供只读查询接口"],
        ["W3", "申请基础信息与入口", "提供关联借测单与纯软件本地申请两种入口", "借测单需客户级联筛选，只允许关联审批通过借测单"],
        ["W3", "方案范围只读锁定", "展示方案编码、名称、关联产品线标签列表 及统一授权时间", "严格只读，杜绝申请人手工增删或替换产品线"],
        ["W3", "业务申请信息", "根据方案特性动态渲染专属业务评估信息（DSOC/加速/组网）", "与底层产品线参数解耦，支持模板化配置与调研文档上传"],
        ["W3", "多产品线工作台", "垂直导航逐项切换产品线参数，支持按需产品线滑动开关", "未上线产品线默认关闭且免校验，开启后纳入整单必填校验"],
        ["W3", "全产品原子授权", "整单进入 W3 审批，审批通过后统一执行授权与失败整单回滚", "任一产品失败整单回滚，不产生部分交付残缺授权"],
        ["W3", "解决方案授权记录", "列表生成一条 SOL 主记录，支持展开查看各产品线明细子行", "主行展示方案与总设备数，展开行展示版本、设备/云图ID及授权码"]
    ]
    create_styled_table(doc, t1_headers, t1_data, col_widths_in=[1.0, 1.8, 2.7, 1.9], header_fill="E5F6FF")

    add_p(doc, "2.2 系统职责", style="Heading 2", sb=10.0, sa=2.0)
    t2_headers = ["系统", "负责"]
    t2_data = [
        ["W3 测试授权系统", "维护解决方案与产品线映射关系及授权天数；向资产系统提供只读配置查询；提供借测单校验与带出；提供多产品线参数工作台；执行 W3 整单审批；调度下游中台原子授权与失败回滚；生成并展示 SOL 汇总与展开明细记录。"],
        ["新资产系统", "读取 W3 启用的解决方案与产品线配置；支持用户在借测单中选择方案并完成借测审批；保存审批快照；向 W3 提供借测单状态与明细查询接口。不主动推送消息至 W3，不自动在 W3 建单。"],
        ["下游许可证中台 / 云图", "接收 W3 下发的产品授权与云服务开通指令，返回具体授权码；提供失败时的批量作废撤销与回滚接口能力。不处理解决方案业务合规审批。"]
    ]
    create_styled_table(doc, t2_headers, t2_data, col_widths_in=[2.2, 5.2], header_fill="E5F6FF")

    add_p(doc, "2.3 待办与单据类型", style="Heading 2", sb=10.0, sa=2.0)
    t3_headers = ["待办/单据类型", "生成时机", "处理人主要动作", "是否审批流"]
    t3_data = [
        ["W3 解决方案申请单", "业务人员在 W3 关联借测单或纯软件模式提交后生成", "申请人填写完整参数并确认整单提交", "否（业务发起）"],
        ["W3 整单审批待办", "解决方案申请单提交成功后进入待办池", "W3 审批人审核方案信息、业务调研与产品配置，给出通过/驳回结论", "是（整单合规审批）"]
    ]
    create_styled_table(doc, t3_headers, t3_data, col_widths_in=[1.5, 2.2, 2.4, 1.3], header_fill="E5F6FF")

    add_p(doc, "2.4 审批与处理口径", style="Heading 2", sb=10.0, sa=2.0)
    t4_headers = ["事项", "口径"]
    t4_data = [
        ["审批单粒度", "整套解决方案作为一张申请单统一审批，不按产品线拆分审批单。"],
        ["审批驳回机制", "任一审批节点驳回时整单驳回，申请人可基于原单修改数据后重新提交。"],
        ["授权交付原子性", "方案内全部产品授权成功方为整单成功；任一产品失败整单标记为失败，并回滚已开通产品。"],
        ["客户云图身份口径", "含 XaaS 产品线的方案，授权对象为客户云图账号下的云图ID，缺失云图身份直接阻断提交。"]
    ]
    create_styled_table(doc, t4_headers, t4_data, col_widths_in=[1.8, 5.6], header_fill="E5F6FF")

    # 3. 业务流程
    add_p(doc, "3. 业务流程", style="Heading 1", sb=24.0, sa=4.0)
    add_p(doc, "3.1 解决方案型授权申请与执行流程", style="Heading 2", sb=10.0, sa=2.0)
    add_body_p(doc, "1、W3 配置管理员在「系统设置」→「解决方案配置」维护方案编码、名称、产品线组合及统一授权天数（authDays）；新资产系统通过只读接口获取启用的方案数据。")
    add_body_p(doc, "2、业务人员在资产系统发起借测申请并独立完成资产审批；审批通过后，流程在资产系统结束，不向 W3 发送业务消息。")
    add_body_p(doc, "3、业务人员主动登录 W3，进入「解决方案型授权」页面，选择“关联借测单”或“纯软件本地申请”。关联借测单时，输入客户筛选该客户名下的合法借测单，点击“查询并关联”。")
    add_body_p(doc, "4、W3 实时校验借测单准入状态：必须为“审批通过且未归还、未撤销”。校验通过后带出借测单号、申请人、客户、期限及审批结论，并在「方案范围」锁定方案及产品线（只读不可改），按配置天数动态计算统一授权到期时间。")
    add_body_p(doc, "5、若为分布式安全运营、全球加速或全球组网方案，系统展开「业务申请信息」区块，用户补充主打场景、驱动力、预期效果、调研文档或客户需求等专属评估信息。")
    add_body_p(doc, "6、进入「授权参数」多产品线工作台，申请人按左侧导航逐个填写产品参数；未上线的按需产品线（如 SASE-SWG）默认关闭，不参与校验；设备线配置版本、设备 ID 及模块卡片；XaaS 线配置云服务参数并校验客户云图身份；所有产品线授权到期日强制对齐统一到期时间。")
    add_body_p(doc, "7、点击“提交 W3 整单审批”，前置校验全单必填项。校验通过后生成一张 W3 申请单进入审批流。审批人查看整单方案并作出通过或驳回结论。")
    add_body_p(doc, "8、审批通过后执行全产品原子授权：预校验通过后并行下发，全部成功则在授权记录中生成一条以 SOL 开头的主记录，并可展开查看所有产品明细；任一产品失败则整单回滚作废已开通产品，整单标记为失败。")

    # 4. 需求明细
    add_p(doc, "4. 解决方案型测试授权需求明细", style="Heading 1", sb=24.0, sa=4.0)

    # 4.1 解决方案配置管理
    add_p(doc, "4.1 解决方案配置管理", style="Heading 2", sb=10.0, sa=2.0)
    add_subpart_header(doc, "1、初始状态")
    add_body_p(doc, "系统由管理员在配置中心统一维护解决方案与产品线清单的基线配置。初始状态下，系统展示已维护的解决方案定义列表，包含方案编码、方案名称、关联产品线集合、统一授权天数及当前启用状态；列表默认按更新时间倒序排列。新资产系统通过接口同步读取处于“启用”状态的方案配置；若系统中未配置任何方案或方案全部处于停用状态，测试授权申请端不可发起解决方案型授权。")
    add_subpart_header(doc, "2、交互逻辑说明")
    add_body_p(doc, "管理员可按方案编码或方案名称进行关键字检索，或按状态（全部/启用/停用）筛选方案；支持新增与编辑解决方案，录入方案编码、方案名称、统一授权天数，并勾选其所包含的产品线范围（支持硬件设备产品线与 XaaS 云服务产品线混排）；支持对已有方案进行一键启用/停用切换，停用后新资产系统与测试授权申请端即时不可选；提供模拟接口读取功能，支持实时查看新资产系统调用的方案配置数据快照。")
    add_subpart_header(doc, "3、字段校验规则")
    add_body_p(doc, "方案编码必填且全局唯一，不可包含特殊字符，新增后锁定只读不可修改；方案名称必填且系统内不可重名；统一授权天数必填，必须为大于 0 的正整数（用于后续申请时统一锁定全方案各产品线的授权周期）；关联产品线必选，至少勾选一条产品线；已产生历史测试授权单据的方案禁止物理删除，仅允许变更为停用状态。")
    add_subpart_header(doc, "4、提交逻辑处理")
    add_body_p(doc, "方案新增、编辑或启停状态变更后即时持久化生效，自动记录管理员操作审计日志（操作人、时间、变更内容）；状态变动实时生效于测试授权申请端下拉选项及新资产系统同步接口；历史已审批完成的授权单据按申请时的历史快照生效，不受后续方案配置调整的影响。")
    add_subpart_header(doc, "5、验收标准")
    add_body_p(doc, "方案列表展示准确，支持多维度检索与启停流转；方案编码与名称防重校验拦截生效；授权天数非正整数、产品线未勾选等异常前置阻断；停用方案在申请端和资产系统中即时不可选；配置变更后不影响历史单据数据完整性。")
    
    t_cfg_headers = ["字段", "数据来源", "可编辑", "必填", "说明"]
    t_cfg_data = [
        ["方案编码", "管理员新增录入", "新增时可填，编辑时只读", "是", "系统全局唯一标识，如 SOL-CLOUD-001"],
        ["方案名称", "管理员录入", "是", "是", "解决方案全称，如企业云一体化方案"],
        ["关联产品线", "勾选产品目录", "是", "是", "至少选择一个产品线，支持设备线与 XaaS 线混排"],
        ["授权时间", "管理员录入天数", "是", "是", "正整数天数，如 30、60、90，用于申请时计算统一到期日"],
        ["状态", "系统维护", "操作列点击切换", "是", "启用 (enabled) / 停用 (disabled)"]
    ]
    create_styled_table(doc, t_cfg_headers, t_cfg_data, col_widths_in=[1.2, 1.5, 1.2, 0.8, 2.7], header_fill="1F4E78", header_text_color="FFFFFF")
    add_caption(doc, "图 1  解决方案配置管理原型参考")

    # 4.2 解决方案授权申请
    add_p(doc, "4.2 解决方案授权申请", style="Heading 2", sb=10.0, sa=2.0)

    # 4.2.1
    add_p(doc, "4.2.1 申请入口与借测单关联", style="Heading 3", sb=10.0, sa=2.0)
    add_subpart_header(doc, "1、初始状态")
    add_body_p(doc, "申请人进入测试授权申请页面，系统提供「关联借测单」与「纯软件本地申请」两种申请模式。初始状态下，系统默认处于「关联借测单」模式，页面处于待选定客户与待选择借测单状态，方案范围及各产品线参数暂不加载；若本次为纯软件方案试用或无需借用实体硬件设备，申请人可主动切换为「纯软件本地申请」模式。")
    add_subpart_header(doc, "2、交互逻辑说明")
    add_body_p(doc, "在关联借测单模式下，申请人先搜索并选定客户，系统实时拉取该客户名下在新资产系统中已审批的所有借测单数据供选择。点击「查询并关联」，系统执行借测单准入强校验：仅“审批通过且未归还、未撤销”且未被其他有效测试授权占用的借测单允许准入。准入通过后带出借测关键信息（借测单号、申请人、客户、期限、产品数），并自动锁定加载方案范围。若借测单处于“审批中/未审批”、“已归还”、“已撤销”或“已被关联占用”，系统弹出明确阻断提示卡片，严禁进入后续环节。在纯软件申请模式下，申请人直接选择已启用的方案，录入接收邮箱（选填）并关联客户，同样自动锁定加载方案范围。方案范围固化锁定：方案包含的产品线清单及统一授权天数（authDays）严格只读锁定，杜绝申请人手工增删或替换产品线；系统依据 authDays 自动联动计算统一授权到期日（申请日 + authDays），并强制分发至方案下所有产品线，保持全方案授权周期严格一致。")
    add_subpart_header(doc, "3、字段校验规则")
    add_body_p(doc, "授权类型必选；关联借测单模式下客户与合法借测单必选；纯软件模式下方案与客户必选；仅“审批通过且未归还、未撤销”借测单可通过准入校验；方案产品组合只读不可篡改；统一授权到期时间不可为空。")
    add_subpart_header(doc, "4、提交逻辑处理")
    add_body_p(doc, "将申请来源标识（borrow / local）、借测单全量元数据快照、锁定的方案编码、名称、产品线列表及统一到期日深拷贝存入单据模型中，随单提交。")
    add_subpart_header(doc, "5、验收标准")
    add_body_p(doc, "两类入口切换顺畅且数据互斥隔离；未审批、已归还、已撤销借测单 100% 阻断拦截；合法借测单回填准确；方案范围与产品清单严格只读锁定；统一授权到期日准确联动至所有产品线。")
    
    t_entry_headers = ["字段", "数据来源", "可编辑", "必填", "说明"]
    t_entry_data = [
        ["授权类型", "下拉选项", "是", "是", "关联借测单 (borrow) / 纯软件本地申请 (local)"],
        ["客户", "主数据客户名录", "是", "是", "输入名称/编码下拉点选，展示 KA/普通徽章与云图账号"],
        ["借测单号", "新资产系统", "是", "关联借测单必选", "选定客户后级联激活，展示单号与借测状态"],
        ["解决方案", "W3 启用方案库", "是", "纯软件申请必选", "格式如 SOL-ACCEL-006 | 全球加速方案-SaaS版"],
        ["接收邮箱", "用户手动输入", "是", "否", "建议优先填客户邮箱，用于接收授权结果通知"]
    ]
    create_styled_table(doc, t_entry_headers, t_entry_data, col_widths_in=[1.2, 1.5, 1.0, 1.2, 2.5], header_fill="1F4E78", header_text_color="FFFFFF")
    add_caption(doc, "图 2  关联借测单与准入状态校验原型参考")

    # 4.2.2
    add_p(doc, "4.2.2 业务信息与产品线参数配置（重点规范 XaaS 与云图 ID 校验）", style="Heading 3", sb=10.0, sa=2.0)
    add_subpart_header(doc, "1、初始状态")
    add_body_p(doc, "方案范围锁定后，系统根据《安全方案测试申请规则.xlsx》中该方案绑定的组件清单与业务模板，初始化「业务申请信息」与「多产品线参数工作台」两部分填报区域。业务信息部分根据方案规则动态呈现专属调研字段（初始处于待填报状态）；产品线工作台初始化该方案下所有产品线的参数填报面板，核心必选产品线默认处于待填报状态，按需/可选产品线默认处于未激活状态。针对 XaaS 云服务产品线（SASE-GA、SASE-SWG、SASE-ZTNA、SaaS-XDR、MSS 等），系统自动隐藏硬件设备字段，并根据已选客户自动检测并带出客户在云图中台绑定的「云图 ID（tenantId）」；若客户尚未在云图中台开户，系统处于未开户阻断状态。")
    add_subpart_header(doc, "2、交互逻辑说明")
    add_body_p(doc, "（1）业务申请信息动态渲染：系统所支持的安全解决方案（包括全球加速方案、全球组网方案、安全托管方案、托管式安全运营方案、分布式安全运营方案、数据安全方案等）以及各方案所需的业务调研信息字段（如测试主打场景、驱动力、预期效果、调研文档、竞对等），统一以《安全方案测试申请规则.xlsx》为基线进行动态渲染配置，申请人无需关心技术底层表单结构。后续新增方案或调整调研字段，通过后台规则配置即时生效，无需开发介入。\n"
               "（2）多产品线按需填报：申请人可在左侧产品线列表中切换当前填报产品，已录入数据平滑保留；针对按需/可选产品线，申请人可点击开关激活，激活后右侧展开该产品参数配置表单并纳入必填校验范围；关闭时免除校验并不参与本次申请。\n"
               "（3）硬件设备产品线参数填报：继承借测单回填的设备 ID 与设备 SN，申请人选定软件版本，并可按需开启或选配功能模块卡片（如 AC/AF 的安全订阅、HCI 的算力存储规格）。\n"
               "（4）XaaS 产品线与云图 ID 强校验交互（重点规范）：\n"
               "  - 完全隐藏硬件字段：XaaS 产品界面彻底隐藏“设备 ID”与“设备 SN”输入项与展示项，杜绝伪硬件信息；\n"
               "  - 云图 ID (tenantId) 为唯一授权主体：系统自动读取并校验客户在深信服云图中台关联的租户账号与云图 ID；\n"
               "  - 未入驻强行阻断：若客户尚未在云图中台开户或缺失有效云图 ID，系统在工作台面板顶部弹出高亮拦截告警：“【阻断】当前方案包含 XaaS 云服务，检测到该客户尚未在云图中台开户或未关联云图 ID。无法发起测试授权申请，请先联系管理员或前往云图中台完成客户租户建档入驻。”，并强行阻断整单提交，严禁绕过；\n"
               "  - XaaS 业务服务规格录入：申请人录入 XaaS 专属规格，包括 SASE-GA 带宽档位（5M/10M/15M/20M/30M/50M）、SASE-ZTNA 授权并发用户数（正整数，默认 50）、SaaS-XDR/MSS 接入日志量与上传客户信息收集表（.xlsx/.xls）等；\n"
               "  - 授权记录数据映射：下发交付时调用云图中台租户授权接口；在授权记录中，XaaS 产品线设备 ID 字段明确展示为“云图ID: [tenantId]”，设备 SN 固定展示为破折号“—”；客户如果存在多个云图租户账号，系统支持下拉选择，并在选择后自动回显该租户的所属行业、开通状态及管理员账号。")
    add_subpart_header(doc, "3、字段校验规则")
    add_body_p(doc, "云图 ID 强校验（核心阻断项）：方案若包含任一 XaaS 产品线，已选客户必须拥有有效合法的云图 ID（tenantId）。无云图身份或云图状态异常时，全单坚决阻断提交，拦截率 100%；业务信息校验：严格对齐《安全方案测试申请规则.xlsx》中各方案必填项（如必选调研文档、必选场景等），未补全阻断提交；硬件产品参数校验：版本必填，借测设备 ID/SN 必须有效，已启用模块参数符合合法范围；XaaS 规格参数校验：SASE-GA 带宽必选；SASE-ZTNA 用户数必须为大于等于 1 的正整数；SaaS-XDR/MSS 日志量与信息收集表必填。")
    add_subpart_header(doc, "4、提交逻辑处理")
    add_body_p(doc, "组装包含业务信息快照与各激活产品线参数快照的数据包。针对 XaaS 产品线，提交数据中的 deviceId 统一固化为“云图ID: [tenantId]”，deviceSn 统一写入破折号“—”；未激活的按需产品线彻底从提交快照中剔除，不生成下游授权任务。")
    add_subpart_header(doc, "5、验收标准")
    add_body_p(doc, "业务信息准确按《安全方案测试申请规则.xlsx》动态展开；硬件产品线与 XaaS 产品线字段清晰隔离，XaaS 界面 100% 隐藏设备 ID 与 SN；无云图账号客户申请包含 XaaS 的方案时，前置告警与提交拦截 100% 触发；具备有效云图 ID 的客户能顺畅填写规格参数并通过校验；按需产品线开关开启/关闭逻辑与数据过滤完全正确。")

    t_xaas_headers = ["字段/对象", "适用产品", "数据来源", "必填", "说明"]
    t_xaas_data = [
        ["云图ID (tenantId)", "所有 XaaS 产品线 (SASE/XDR/MSS等)", "云图中台租户身份绑定", "是 (核心阻断项)", "XaaS 授权唯一核心标识；客户无云图账号或无云图ID时系统强行阻断提交并引导建档入驻"],
        ["带宽档位", "SASE-GA", "下拉选项", "是", "档位可选 5M / 10M / 15M / 20M / 30M / 50M，满足各加速测试场景"],
        ["授权用户数", "SASE-ZTNA / aTrust", "数字输入框", "是", "并发接入用户数，必须为正整数（如 50、100）"],
        ["日志量 / 收集表", "SaaS-XDR / MSS", "文本输入 / Excel文件上传", "是", "单位为万条/天，客户现场信息收集表必传 (.xlsx/.xls)"],
        ["是否 PK / 友商名", "SaaS-XDR / XDR", "下拉选项 / 文本输入", "是", "选择存在竞对 PK 时动态展开友商名称录入框，必填"]
    ]
    create_styled_table(doc, t_xaas_headers, t_xaas_data, col_widths_in=[1.5, 1.5, 1.5, 0.8, 2.1], header_fill="1F4E78", header_text_color="FFFFFF")
    add_caption(doc, "图 3  多产品线工作台与 XaaS 云图身份校验原型参考")

    # 4.2.3
    add_p(doc, "4.2.3 解决方案整单审批与原子交付回滚", style="Heading 3", sb=10.0, sa=2.0)
    add_subpart_header(doc, "1、初始状态")
    add_body_p(doc, "申请单在所有已激活产品线参数校验通过前处于编辑草稿状态，审批流程未发起。系统处于整单校验前置监听状态。")
    add_subpart_header(doc, "2、交互逻辑说明")
    add_body_p(doc, "点击提交按钮，顺序执行严格前置校验（单据与客户有效性、业务信息完整性、XaaS 云图身份有效性、所有激活产品线必填参数与模块卡片完整性）。任一不符在页面高亮报错并阻断。\n"
               "校验通过后生成统一的 W3 方案申请单号（SOL-...）进入待审批状态。整单走统一审批流，任一审批节点驳回则整单驳回，申请人可修改重提。\n"
               "审批通过后触发下游授权原子下发：系统调用各产品线发号中台与云图中台开通授权。若所有产品均成功，整单标记为“已授权”并归档；若任一产品线开通失败，立即触发整单事务回滚机制，自动调用作废接口撤销已开通产品的 License，整单标记为“开通失败”，杜绝方案残缺交付。")
    add_subpart_header(doc, "3、字段校验规则")
    add_body_p(doc, "前置校验覆盖全单所有激活产品线，不允许存在任何空缺必填项；终审执行前再次对齐借测单状态。")
    add_subpart_header(doc, "4、提交逻辑处理")
    add_body_p(doc, "生成包含全套参数深复制快照的申请单据，维系统一全局事务 ID 用于回滚追踪。")
    add_subpart_header(doc, "5、验收标准")
    add_body_p(doc, "任意漏填必填项点击提交均能精准定位提示；提交后单据状态流转正确；失败整单回滚机制彻底杜绝残缺授权。")

    # 4.3 授权记录展示与明细检索
    add_p(doc, "4.3 授权记录展示与明细检索", style="Heading 2", sb=10.0, sa=2.0)
    add_subpart_header(doc, "1、初始状态")
    add_body_p(doc, "审批通过并完成授权下发后，统一授权记录列表中以方案为维度生成一条汇总主记录（SOL- 前缀），汇聚方案整体状态、客户信息、统一到期日及所含产品线总览，各产品线独立明细默认折叠归属在主记录下。")
    add_subpart_header(doc, "2、交互逻辑说明")
    add_body_p(doc, "主列表支持在类型筛选下拉中选择“解决方案”，精准过滤出所有解决方案主记录；输入设备 ID、云图 ID 或 SN 搜索均能命中并展示所属方案主行。点击主记录行或展开图标，可展开查看嵌套的明细子表格；再次点击收起。\n"
               "明细子表格中：设备产品线展示真实设备 ID 与序列号；XaaS 产品线在“设备ID/云图ID”列展示“云图ID: [tenantId]”，在“设备SN”列展示破折号“—”。各产品行均展示方案统一的授权有效期与中台生成的真实授权码（如 LIC-HCI-806、LIC-SaaS-XDR-806）。")
    add_subpart_header(doc, "3、字段校验规则")
    add_body_p(doc, "主行设备总数必须与实际启用产品线数严格一致；主行有效期起止范围必须与方案统一到期日吻合；展开子行数量必须与申请时激活的产品线数完全相等；XaaS 产品线严禁展示伪设备 SN。")
    add_subpart_header(doc, "4、提交逻辑处理")
    add_body_p(doc, "数据层以一对多主子表架构存储，主行承载整单状态与统计指标，子表承载各产品 License 明细，随整单原子授权成功后一次性写入并绑定主行编号。")
    add_subpart_header(doc, "5、验收标准")
    add_body_p(doc, "解决方案展示为单条主行；主行指标统计准确；搜索与类型筛选联动精确；展开子表格排版整齐无错位；设备与 XaaS 产品线字段区分展示准确；所有产品授权码完整呈现。")

    t_rec_headers = ["列名", "适用产品类型", "展示内容口径", "说明"]
    t_rec_data = [
        ["产品线", "所有产品", "产品线名称（如 HCI、AF、SaaS-XDR）", "方案下各产品独立成行"],
        ["版本号", "所有产品", "申请时确定的具体版本", "如 HCI-6.8.1-VKEY、SaaS-XDR 标准版"],
        ["设备ID / 云图ID", "设备线 / XaaS线", "设备线显示真实设备ID；XaaS显示云图ID", "XaaS 口径如 云图ID: YT-10010101"],
        ["设备SN", "设备线 / XaaS线", "设备线显示序列号；XaaS显示破折号 —", "XaaS 绝不展示虚假硬件序列号"],
        ["授权有效期", "所有产品", "至 [统一到期日]", "方案下所有产品严格统一到期时间"],
        ["授权码", "所有产品", "下游中台生成的 License 凭证码", "格式如 LIC-[产品线]-806"]
    ]
    create_styled_table(doc, t_rec_headers, t_rec_data, col_widths_in=[1.4, 1.4, 2.3, 2.3], header_fill="1F4E78", header_text_color="FFFFFF")
    add_caption(doc, "图 4  解决方案授权记录汇总与展开明细原型参考")

    # 5. 外部协同与接口需求明细
    add_p(doc, "5. 外部协同与接口需求明细", style="Heading 1", sb=24.0, sa=4.0)
    add_p(doc, "5.1 新资产系统协同", style="Heading 2", sb=10.0, sa=2.0)
    add_body_p(doc, "5.1.1 W3 解决方案配置查询接口：新资产系统通过该接口获取 W3 处于启用状态的解决方案编码、名称、产品线类型列表（区分 device/xaas）及授权天数（authDays）。新资产系统仅在用户建借测单时读取，不嵌入 W3 配置页。")
    add_body_p(doc, "5.1.2 借测单状态与快照查询接口：W3 申请人在关联借测单时调用新资产系统，获取借测单号、审批状态、客户、借测人、借测起止日、归还状态、撤销状态及已审批方案范围快照；若含 XaaS 线同步返回客户绑定的云图账号与云图ID。")
    
    add_p(doc, "5.2 下游许可证平台与云图中台协同", style="Heading 2", sb=10.0, sa=2.0)
    add_body_p(doc, "5.2.1 批量授权下发接口：整单审批通过后，W3 向下游中台派发批量开通任务，按各产品线参数生成 License 文件或激活云端租户实例。")
    add_body_p(doc, "5.2.2 异常原子回滚接口：两阶段事务保障。任一产品开通失败，W3 立即调用中台作废接口撤销已开通的产品凭证，确保整套方案交付结果的一致性。")

    # 6. 异常与一致性处理矩阵
    add_p(doc, "6. 异常与一致性处理矩阵", style="Heading 1", sb=24.0, sa=4.0)
    t_ex_headers = ["异常场景", "触发时机", "系统判定与处理机制", "申请人提示文案与引导"]
    t_ex_data = [
        ["借测单号不存在", "点击查询并关联", "资产系统检索无此单号", "页面提示：未找到该借测单，请核对单号后重试。"],
        ["借测单未审批通过", "点击查询并关联", "approvalStatus !== 审批通过", "红色错误卡片：无法关联：该借测单尚未审批通过，当前状态为[状态]。"],
        ["借测设备已归还", "点击查询并关联", "returned === true", "红色错误卡片：无法关联：该借测单已归还。"],
        ["借测单已撤销", "点击查询并关联", "revoked === true", "红色错误卡片：无法关联：该借测单已撤销。"],
        ["资产接口异常/超时", "点击查询并关联", "网络超时或接口 500", "红色提示：查询失败：资产系统暂时不可用，请稍后重试。保留已填数据。"],
        ["方案含 XaaS 缺云图身份", "提交整单审批", "客户名下无可用 tenantId", "提交栏报错：XaaS 产品线需选择已关联云图账号和云图ID的客户。阻断提交。"],
        ["DSOC 缺调研文档", "提交整单审批", "调研文档上传项为空", "提交栏报错：请补全业务申请信息（分布式安全运营）。阻断提交。"],
        ["常规产品线参数不全", "提交整单审批", "某激活产品漏填版本/规格", "提交栏报错：[产品线名称] 的授权参数未填写完整。精准定位至该产品。"],
        ["方案配置已被管理员停用", "纯软件本地申请", "方案状态为 disabled", "本地申请下拉列表中直接剔除，不可选；历史已审批借测单仍允许关联。"],
        ["单个产品授权下发失败", "审批通过下发时", "中台返回开通超时或失败", "立即触发整单回滚，作废已开通产品，整单标记“整单失败”，杜绝残缺交付。"],
        ["中台回滚作废异常", "执行整单回滚时", "网络抖动导致撤销挂起", "记录告警并派发后台人工补偿工单，对外保持不可交付状态。"],
        ["借测单方案与当前不同", "点击查询并关联", "资产已审批方案产品线被调整", "以资产已审批快照为准只读带出，不按 W3 最新配置改写历史审批内容。"]
    ]
    create_styled_table(doc, t_ex_headers, t_ex_data, col_widths_in=[1.5, 1.4, 2.3, 2.2], header_fill="E5F6FF")

    # 7. 业务方需求落地：业务申请信息与产品组合灵活配置规范
    add_p(doc, "7. 业务方需求落地：业务申请信息与产品组合灵活配置规范", style="Heading 1", sb=24.0, sa=4.0)
    add_p(doc, "7.1 架构设计背景与免代码开发原则", style="Heading 2", sb=10.0, sa=2.0)
    add_body_p(doc, "业务方明确要求：“页面的业务申请信息、产品授权组合要考虑灵活配置。后续随方案推广，需要收集的信息和方案产品组合可能会有变化，不要每次调整都要开发”。")
    add_body_p(doc, "针对该核心诉求，系统在架构上确立免代码开发扩展原则：底层实现方案元数据驱动模型（Metadata-Driven Model），将产品授权组合与业务调研表单与前端静态页面彻底解耦。推广新方案或微调调研字段时，无需进行全链路代码重新发版。")

    add_p(doc, "7.2 产品组合元数据驱动方案", style="Heading 2", sb=10.0, sa=2.0)
    add_body_p(doc, "方案与产品线的组合关系完全由 W3「解决方案配置」数据表驱动。管理员在抽屉中自由勾选产品目录并配置授权天数后，申请端根据接口返回的产品线动态构建左侧工作台导航，复用各产品的标准参数组件，产品组合调整 0 代码发版。")

    add_p(doc, "7.3 业务表单模板字典与扩展机制", style="Heading 2", sb=10.0, sa=2.0)
    add_body_p(doc, "业务申请信息采用模板化解析引擎。系统预置业务模板字典：NONE（通用无专属字段）、DSOC_TEMPLATE（安全运营类）、ACCEL_TEMPLATE（加速类）、NETWORK_TEMPLATE（组网类）。模板以 JSON Schema 定义字段标识、控件类型、选项源、校验规则及条件显隐。新增方案或增减调研字段时，仅需在后台维护模板映射，前端引擎自适应渲染并自适应审批列表展示。")
    
    t_tpl_headers = ["模板标识", "适用方案", "核心字段与组件", "扩展方式"]
    t_tpl_data = [
        ["NONE", "通用解决方案（如企业云）", "无特殊业务信息，直接进入工作台", "标准开箱即用，无需配置"],
        ["DSOC_TEMPLATE", "分布式安全运营方案", "主打场景、驱动力、预期效果、调研文档上传、转销计划", "支持在模板 Schema 中增减多选题项及附件类型"],
        ["ACCEL_TEMPLATE", "全球加速方案（SaaS/本地版）", "客户需求、项目规模、竞争对手、所在省份、备注", "通过 condition 字段动态控制字段在不同子方案下的显隐"],
        ["NETWORK_TEMPLATE", "全球组网方案", "客户需求、项目规模、竞争对手（收起省份/备注）", "复用加速模板核心字段，自动屏蔽地域属性"],
        ["CUSTOM_TEMPLATE", "未来推广新方案", "按业务调研诉求自由组合文本、下拉、复选与文件控件", "直接新增模板配置，前端自适应渲染，免发版开发"]
    ]
    create_styled_table(doc, t_tpl_headers, t_tpl_data, col_widths_in=[1.5, 1.8, 2.5, 1.6], header_fill="E5F6FF")

    # 8. 待开发确认项
    add_p(doc, "8. 待开发确认项", style="Heading 1", sb=24.0, sa=4.0)
    add_body_p(doc, "1. 确认 W3 整单审批的正式节点数量、审批人角色映射及超时自动审批规则。")
    add_body_p(doc, "2. 确认新资产系统与 W3 配置同步接口的认证鉴权方式、心跳刷新频率与超时熔断策略。")
    add_body_p(doc, "3. 确认新资产系统借测单撤销或设备归还后，W3 对已下发有效授权的自动作废与联动清理机制。")
    add_body_p(doc, "4. 确认下游许可证平台针对各产品线批量撤销/作废接口的事务一致性与幂等性保障能力。")
    add_body_p(doc, "5. 确认多云图账号客户在前端申请时的下拉点选交互原型与默认关联规则。")
    add_body_p(doc, "6. 确认业务申请信息 JSON Schema 模板存储位置（数据库表 vs 配置中心）与后台维护界面的上线排期。")

    print(f"Saving exact rebuild docx to {OUTPUT_DOCX_PATH}...")
    doc.save(str(OUTPUT_DOCX_PATH))
    print(f"Success! Rebuilt docx size: {os.path.getsize(OUTPUT_DOCX_PATH)} bytes")

if __name__ == "__main__":
    main()
